"""Document parser — abstraction over document-to-text conversion.

Uses PyMuPDF for PDF parsing, python-docx for DOCX, and direct reading
for plain text and Markdown files.
"""

import logging
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)

# Supported file extensions
SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}


@dataclass
class ParsedDocument:
    """Result of parsing a document."""
    text: str
    pages: list[str] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)

    @property
    def page_count(self) -> int:
        return len(self.pages) if self.pages else 1

    @property
    def is_empty(self) -> bool:
        return not self.text or not self.text.strip()


class DocumentParser:
    """Parse documents into structured text.

    Uses lightweight libraries:
    - PyMuPDF (fitz) for PDF — fast, no ML dependencies
    - python-docx for DOCX
    - Direct read for TXT/MD

    The ingestion service calls parser.parse(file_path) and receives
    a ParsedDocument regardless of format.
    """

    def parse(self, file_path: Path) -> ParsedDocument:
        """Parse a document file into structured text.

        Args:
            file_path: Path to the document file.

        Returns:
            ParsedDocument with extracted text and metadata.

        Raises:
            ValueError: If the file type is not supported.
            FileNotFoundError: If the file does not exist.
            RuntimeError: If parsing fails.
        """
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        suffix = file_path.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {suffix}. "
                f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )

        logger.info("Parsing document: %s (type: %s)", file_path.name, suffix)

        parsers = {
            ".txt": self._parse_text_file,
            ".md": self._parse_text_file,
            ".pdf": self._parse_pdf,
            ".docx": self._parse_docx,
        }
        return parsers[suffix](file_path)

    def _parse_text_file(self, file_path: Path) -> ParsedDocument:
        """Parse plain text or markdown files by direct reading."""
        text = file_path.read_text(encoding="utf-8")

        if not text.strip():
            raise RuntimeError(f"Document is empty: {file_path.name}")

        return ParsedDocument(
            text=text,
            pages=[text],
            metadata={
                "source": file_path.name,
                "file_type": file_path.suffix.lower(),
                "parser": "direct",
            },
        )

    def _parse_pdf(self, file_path: Path) -> ParsedDocument:
        """Parse PDF using PyMuPDF (fitz)."""
        try:
            import fitz  # PyMuPDF
        except ImportError:
            raise RuntimeError(
                "PyMuPDF is required for PDF parsing. Install: pip install pymupdf"
            )

        try:
            doc = fitz.open(str(file_path))
            pages: list[str] = []
            full_text_parts: list[str] = []

            for page_num in range(len(doc)):
                page = doc[page_num]
                page_text = page.get_text("text")
                pages.append(page_text)
                full_text_parts.append(page_text)

            doc.close()
            full_text = "\n\n".join(full_text_parts)

            if not full_text.strip():
                raise RuntimeError(f"Document is empty after parsing: {file_path.name}")

            return ParsedDocument(
                text=full_text,
                pages=pages,
                metadata={
                    "source": file_path.name,
                    "file_type": ".pdf",
                    "parser": "pymupdf",
                    "page_count": len(pages),
                },
            )
        except RuntimeError:
            raise
        except Exception as e:
            raise RuntimeError(f"Failed to parse PDF {file_path.name}: {e}") from e

    def _parse_docx(self, file_path: Path) -> ParsedDocument:
        """Parse DOCX using python-docx."""
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError(
                "python-docx is required for DOCX parsing. Install: pip install python-docx"
            )

        try:
            doc = Document(str(file_path))
            paragraphs: list[str] = []

            for para in doc.paragraphs:
                if para.text.strip():
                    # Preserve heading structure as Markdown
                    if para.style and para.style.name.startswith("Heading"):
                        level = para.style.name.replace("Heading ", "")
                        try:
                            hashes = "#" * int(level)
                        except ValueError:
                            hashes = "##"
                        paragraphs.append(f"{hashes} {para.text}")
                    else:
                        paragraphs.append(para.text)

            full_text = "\n\n".join(paragraphs)

            if not full_text.strip():
                raise RuntimeError(f"Document is empty after parsing: {file_path.name}")

            return ParsedDocument(
                text=full_text,
                pages=[full_text],
                metadata={
                    "source": file_path.name,
                    "file_type": ".docx",
                    "parser": "python-docx",
                },
            )
        except RuntimeError:
            raise
        except Exception as e:
            raise RuntimeError(f"Failed to parse DOCX {file_path.name}: {e}") from e
